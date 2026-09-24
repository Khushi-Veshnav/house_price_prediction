"""Compile executive reports (Word + HTML) from model/metrics.json and report_images/.

Usage:  python generate_report.py
Any extra PNGs named ui_*.png in report_images/ (e.g. Streamlit screenshots) are embedded too.
"""
import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Inches

ROOT = Path(__file__).parent
IMG = ROOT / "report_images"
m = json.loads((ROOT / "model" / "metrics.json").read_text())
charts = [
    ("price_vs_area.png", "Price vs area, coloured by number of rooms"),
    ("actual_vs_predicted.png", "Actual vs predicted prices on the held-out test set"),
    ("residuals.png", "Distribution of residuals"),
    ("model_comparison.png", "Test R² of Linear, Ridge and Lasso"),
]
charts += [(p.name, f"Application screenshot: {p.stem}") for p in sorted(IMG.glob("ui_*.png"))]
charts = [(f, c) for f, c in charts if (IMG / f).exists()]

best = m["results"][m["model"]]
summary = (
    f"A {m['model']} model trained on {m['n_train']} houses and evaluated on {m['n_test']} unseen houses "
    f"explains {best['R2']:.1%} of price variance (MAE ${best['MAE']:,.0f}, RMSE ${best['RMSE']:,.0f}). "
    "Ridge and Lasso perform almost identically, so the simpler unregularised model is deployed."
)
coef_lines = [f"{k}: {v:,.0f} per standard deviation" for k, v in m["coefficients"].items()]

# ---------------- Word ----------------
doc = Document()
doc.add_heading("House Price Prediction – Executive Report", 0)
doc.add_paragraph(f"Generated {date.today():%d %B %Y}")
doc.add_heading("Executive summary", 1)
doc.add_paragraph(summary)
doc.add_heading("Approach", 1)
for t in ["Data: house area, number of rooms and sale price.",
          "Preprocessing: StandardScaler fitted on the training split only.",
          "Models compared: Linear Regression, Ridge, Lasso (80/20 split, 5-fold CV).",
          "Serving: Flask REST API (/predict) with a Streamlit front end."]:
    doc.add_paragraph(t, style="List Bullet")
doc.add_heading("Results", 1)
table = doc.add_table(rows=1, cols=5)
table.style = "Light Grid Accent 1"
for i, h in enumerate(["Model", "MAE", "RMSE", "Test R²", "CV R²"]):
    table.rows[0].cells[i].text = h
for name, r in m["results"].items():
    row = table.add_row().cells
    row[0].text = name
    row[1].text = f"{r['MAE']:,.0f}"
    row[2].text = f"{r['RMSE']:,.0f}"
    row[3].text = f"{r['R2']:.3f}"
    row[4].text = f"{r['CV_R2']:.3f}"
doc.add_heading("Model coefficients (scaled features)", 1)
for line in coef_lines:
    doc.add_paragraph(line, style="List Bullet")
doc.add_paragraph(f"Intercept: {m['intercept']:,.0f}")
doc.add_heading("Charts", 1)
for f, cap in charts:
    doc.add_picture(str(IMG / f), width=Inches(5.2))
    doc.add_paragraph().add_run(cap).italic = True
doc.save(ROOT / "House_Price_Prediction_Report.docx")

# ---------------- HTML ----------------
import base64
rows = "".join(
    f"<tr><td>{n}</td><td>{r['MAE']:,.0f}</td><td>{r['RMSE']:,.0f}</td><td>{r['R2']:.3f}</td><td>{r['CV_R2']:.3f}</td></tr>"
    for n, r in m["results"].items())
figs = "".join(
    f"<figure><img src='data:image/png;base64,{base64.b64encode((IMG / f).read_bytes()).decode()}'/>"
    f"<figcaption>{c}</figcaption></figure>" for f, c in charts)
html = f"""<!doctype html><html><head><meta charset="utf-8"><title>House Price Prediction Report</title>
<style>body{{font-family:system-ui,sans-serif;max-width:860px;margin:2rem auto;padding:0 1rem;color:#222}}
table{{border-collapse:collapse;width:100%}}td,th{{border:1px solid #ccc;padding:6px 10px;text-align:right}}
td:first-child,th:first-child{{text-align:left}}th{{background:#f0f4f8}}img{{max-width:100%}}
figure{{margin:1.5rem 0}}figcaption{{font-style:italic;color:#555}}</style></head><body>
<h1>House Price Prediction – Executive Report</h1><p>Generated {date.today():%d %B %Y}</p>
<h2>Executive summary</h2><p>{summary}</p>
<h2>Results</h2><table><tr><th>Model</th><th>MAE</th><th>RMSE</th><th>Test R²</th><th>CV R²</th></tr>{rows}</table>
<h2>Coefficients (scaled features)</h2><ul>{''.join(f'<li>{l}</li>' for l in coef_lines)}</ul>
<h2>Charts</h2>{figs}</body></html>"""
(ROOT / "house-price-prediction-project-report.html").write_text(html, encoding="utf-8")
print("Wrote House_Price_Prediction_Report.docx and house-price-prediction-project-report.html")
